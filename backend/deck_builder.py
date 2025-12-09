# backend/deck_builder.py
from __future__ import annotations
from pathlib import Path
from typing import Dict, Optional

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from .models import TemplateProfile, GeneratedDeck, GeneratedSlide


def _layout_by_id(prs, profile: TemplateProfile) -> Dict[str, int]:
    return {l.layout_id: l.pptx_layout_index for l in profile.slide_layouts}


def _get_title_placeholder(slide):
    """
    Return the placeholder best suited for title text.
    """
    for shape in slide.placeholders:
        if not shape.is_placeholder:
            continue
        if shape.placeholder_format.type == PP_PLACEHOLDER.TITLE:
            return shape
    # fallback: first text placeholder
    for shape in slide.placeholders:
        if shape.has_text_frame:
            return shape
    return None


def _get_body_placeholder(slide):
    """
    More robust body placeholder selection:
    - Prefer BODY
    - Else largest text-capable placeholder
    """
    candidates = []

    for shape in slide.placeholders:
        if not shape.is_placeholder:
            continue

        pht = shape.placeholder_format.type

        # BODY gets highest priority
        if pht == PP_PLACEHOLDER.BODY:
            return shape

        # If shape can contain text, consider as candidate
        if shape.has_text_frame:
            # Score by bounding box area (largest = main content)
            try:
                area = float(shape.width * shape.height)
            except Exception:
                area = 0
            candidates.append((area, shape))

    if not candidates:
        return None

    # pick largest area
    candidates.sort(key=lambda x: x[0], reverse=True)
    return candidates[0][1]


def _safe_set_textframe(text_frame, paragraphs=None, bullets=None):
    """
    Bullet-safe and crash-safe way to populate a text frame.
    """

    # Reset text frame safely
    text_frame.clear()
    text_frame.text = ""

    if bullets:
        for i, bullet in enumerate(bullets):
            if i == 0:
                text_frame.text = bullet
                text_frame.paragraphs[0].level = 0
            else:
                p = text_frame.add_paragraph()
                p.text = bullet
                p.level = 0

    elif paragraphs:
        for i, para in enumerate(paragraphs):
            if i == 0:
                text_frame.text = para
            else:
                p = text_frame.add_paragraph()
                p.text = para


def build_pptx_from_generated_deck(
    template_path: Path,
    profile: TemplateProfile,
    deck: GeneratedDeck,
    output_path: Path,
) -> Path:

    prs = Presentation(template_path)
    layout_map = _layout_by_id(prs, profile)

    # Clear template sample slides
    while len(prs.slides) > 0:
        r_id = prs.slides._sldIdLst[0].rId
        prs.part.drop_rel(r_id)
        del prs.slides._sldIdLst[0]

    # Build slides
    for slide_data in sorted(deck.slides, key=lambda s: s.slide_index):

        layout_idx = layout_map.get(slide_data.layout_id, None)
        if layout_idx is None:
            layout_idx = 0  # fallback

        layout = prs.slide_layouts[layout_idx]
        slide = prs.slides.add_slide(layout)

        # TITLE placement
        title_shape = _get_title_placeholder(slide)
        if title_shape is not None:
            try:
                title_shape.text = slide_data.title
            except Exception:
                pass

        # BODY placement
        body_shape = _get_body_placeholder(slide)

        if body_shape is not None:
            try:
                tf = body_shape.text_frame
                _safe_set_textframe(
                    tf,
                    paragraphs=slide_data.body_paragraphs,
                    bullets=slide_data.body_bullets
                )
            except Exception as e:
                print(f"[WARNING] Failed to add body text: {e}")

        # Speaker notes
        if slide_data.speaker_notes:
            try:
                notes_slide = slide.notes_slide
                notes_ph = notes_slide.notes_placeholder
                notes_ph.text = slide_data.speaker_notes
            except Exception as e:
                print(f"[WARNING] Could not set speaker notes for slide {slide_data.slide_index}: {e}")

    prs.save(output_path)
    return output_path
